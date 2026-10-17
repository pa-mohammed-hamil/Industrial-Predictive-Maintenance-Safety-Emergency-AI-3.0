"""
scripts/ingest_documents.py

Industrial Predictive Maintenance, Safety &
Emergency Intelligence AI

Document Ingestion Pipeline
===========================

Python 3.11+

Install:
    pip install pypdf python-docx beautifulsoup4 \
        pandas openpyxl sentence-transformers \
        qdrant-client pydantic

Optional OCR:
    pip install pytesseract pdf2image pillow

    OCR also requires Tesseract and Poppler
    to be installed on the operating system.

Supported documents:
    PDF, DOCX, TXT, MD, HTML, CSV, XLSX, JSON

Features:
    - Recursive document discovery
    - Safe extension validation
    - File size restrictions
    - SHA-256 document identification
    - PDF page-level extraction
    - Optional scanned-PDF OCR
    - DOCX paragraph/table extraction
    - Spreadsheet ingestion
    - Text normalization
    - Token-aware chunking
    - Chunk overlap
    - Document and chunk metadata
    - JSONL exports
    - Optional embedding generation
    - Optional Qdrant indexing
    - Tenant-aware payload metadata
    - Structured ingestion reports
    - CLI support

Security:
    - Never executes document content
    - Rejects symbolic links
    - Rejects unsupported file types
    - Limits input size and extracted text
    - Stores tenant metadata for filtering
    - Does not treat retrieved text as instructions

Note:
    Payload tenant_id is not authorization.
    Retrieval services must enforce tenant and
    document permissions independently.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import logging
import re
import sys
import unicodedata

from dataclasses import asdict, dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable
from uuid import UUID, uuid5

from pydantic import BaseModel, ConfigDict, Field


LOGGER = logging.getLogger("document_ingestion")

SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".txt",
    ".md",
    ".markdown",
    ".html",
    ".htm",
    ".csv",
    ".xlsx",
    ".json",
}

CHUNK_NAMESPACE = UUID(
    "b97e9b65-779d-4b1c-8d86-6a6df61fbb47"
)


# ============================================================
# CONFIGURATION
# ============================================================

@dataclass(frozen=True)
class IngestionConfig:
    input_dir: str = "data/documents"

    output_dir: str = "data/processed/documents"

    tenant_id: str = "default"

    chunk_size: int = 450
    chunk_overlap: int = 75

    max_file_size_mb: int = 25
    max_extracted_characters: int = 2_000_000

    enable_ocr: bool = False
    enable_embeddings: bool = False
    enable_qdrant: bool = False

    embedding_model: str = (
        "sentence-transformers/"
        "all-MiniLM-L6-v2"
    )

    qdrant_url: str = "http://localhost:6333"

    qdrant_collection: str = (
        "industrial_documents"
    )

    batch_size: int = 32

    recursive: bool = True

    allowed_extensions: tuple[str, ...] = (
        ".pdf",
        ".docx",
        ".txt",
        ".md",
        ".markdown",
        ".html",
        ".htm",
        ".csv",
        ".xlsx",
        ".json",
    )


# ============================================================
# DATA MODELS
# ============================================================

class DocumentMetadata(BaseModel):
    model_config = ConfigDict(
        extra="forbid"
    )

    document_id: str
    tenant_id: str

    filename: str
    source_path: str
    extension: str

    file_size_bytes: int
    sha256: str

    title: str

    equipment_type: str | None = None
    manufacturer: str | None = None
    asset_id: str | None = None

    document_type: str = "general"

    created_at: str

    extra: dict[str, Any] = Field(
        default_factory=dict
    )


class DocumentSection(BaseModel):
    model_config = ConfigDict(
        extra="forbid"
    )

    section_id: str
    document_id: str

    page_number: int | None = None

    section_title: str | None = None

    text: str

    metadata: dict[str, Any] = Field(
        default_factory=dict
    )


class DocumentChunk(BaseModel):
    model_config = ConfigDict(
        extra="forbid"
    )

    chunk_id: str
    document_id: str
    tenant_id: str

    chunk_index: int

    text: str

    token_count: int

    page_number: int | None = None

    source_path: str

    document_type: str

    equipment_type: str | None = None

    metadata: dict[str, Any] = Field(
        default_factory=dict
    )


class IngestionResult(BaseModel):
    model_config = ConfigDict(
        extra="forbid"
    )

    document_id: str
    filename: str

    status: str

    section_count: int = 0
    chunk_count: int = 0

    error: str | None = None


# ============================================================
# UTILITY FUNCTIONS
# ============================================================

def utc_now() -> str:
    return datetime.now(
        timezone.utc
    ).isoformat()


def sha256_file(
    path: Path,
) -> str:
    digest = hashlib.sha256()

    with path.open("rb") as stream:
        while True:
            block = stream.read(
                1024 * 1024
            )

            if not block:
                break

            digest.update(block)

    return digest.hexdigest()


def stable_id(
    *values: str,
) -> str:
    return str(
        uuid5(
            CHUNK_NAMESPACE,
            "|".join(values),
        )
    )


def normalize_text(
    text: str,
) -> str:
    text = unicodedata.normalize(
        "NFKC",
        text,
    )

    text = text.replace(
        "\x00",
        "",
    )

    text = text.replace(
        "\r\n",
        "\n",
    ).replace(
        "\r",
        "\n",
    )

    text = re.sub(
        r"[ \t]+",
        " ",
        text,
    )

    text = re.sub(
        r"\n{3,}",
        "\n\n",
        text,
    )

    return text.strip()


def infer_document_type(
    filename: str,
) -> str:
    name = filename.lower()

    categories = {
        "equipment_manual": [
            "manual",
            "handbook",
            "datasheet",
            "specification",
        ],
        "maintenance_procedure": [
            "maintenance",
            "service",
            "repair",
            "troubleshooting",
        ],
        "safety_policy": [
            "safety",
            "ppe",
            "hazard",
            "hse",
            "emergency",
        ],
        "inspection_report": [
            "inspection",
            "audit",
            "checklist",
        ],
        "operating_procedure": [
            "sop",
            "operation",
            "procedure",
        ],
        "incident_report": [
            "incident",
            "accident",
            "failure_report",
        ],
    }

    for category, keywords in categories.items():
        if any(
            keyword in name
            for keyword in keywords
        ):
            return category

    return "general"


def infer_equipment_type(
    filename: str,
) -> str | None:
    name = filename.lower()

    equipment_types = [
        "pump",
        "motor",
        "compressor",
        "conveyor",
        "turbine",
        "generator",
        "robot",
        "cnc",
        "boiler",
        "hvac",
        "transformer",
        "valve",
        "bearing",
    ]

    for equipment in equipment_types:
        if equipment in name:
            return equipment

    return None


# ============================================================
# FILE DISCOVERY
# ============================================================

class DocumentDiscovery:
    def __init__(
        self,
        config: IngestionConfig,
    ) -> None:
        self.config = config

    def discover(self) -> list[Path]:
        root = Path(
            self.config.input_dir
        )

        if not root.exists():
            raise FileNotFoundError(
                f"Input directory not found: {root}"
            )

        if not root.is_dir():
            raise ValueError(
                "Input path must be a directory"
            )

        iterator = (
            root.rglob("*")
            if self.config.recursive
            else root.glob("*")
        )

        paths = []

        max_bytes = (
            self.config.max_file_size_mb
            * 1024
            * 1024
        )

        allowed = set(
            self.config.allowed_extensions
        ) & SUPPORTED_EXTENSIONS

        for path in iterator:
            if path.is_symlink():
                LOGGER.warning(
                    "Skipping symbolic link: %s",
                    path,
                )
                continue

            if not path.is_file():
                continue

            if path.suffix.lower() not in allowed:
                continue

            if path.stat().st_size > max_bytes:
                LOGGER.warning(
                    "Skipping oversized file: %s",
                    path,
                )
                continue

            paths.append(path)

        return sorted(paths)


# ============================================================
# DOCUMENT PARSER
# ============================================================

class DocumentParser:
    def __init__(
        self,
        config: IngestionConfig,
    ) -> None:
        self.config = config

    def parse(
        self,
        path: Path,
        document_id: str,
    ) -> list[DocumentSection]:
        extension = path.suffix.lower()

        if extension == ".pdf":
            return self._parse_pdf(
                path,
                document_id,
            )

        if extension == ".docx":
            return self._parse_docx(
                path,
                document_id,
            )

        if extension in {
            ".txt",
            ".md",
            ".markdown",
        }:
            return self._parse_text(
                path,
                document_id,
            )

        if extension in {
            ".html",
            ".htm",
        }:
            return self._parse_html(
                path,
                document_id,
            )

        if extension == ".csv":
            return self._parse_csv(
                path,
                document_id,
            )

        if extension == ".xlsx":
            return self._parse_xlsx(
                path,
                document_id,
            )

        if extension == ".json":
            return self._parse_json(
                path,
                document_id,
            )

        raise ValueError(
            f"Unsupported extension: {extension}"
        )

    def _section(
        self,
        document_id: str,
        index: int,
        text: str,
        page_number: int | None = None,
        title: str | None = None,
        metadata: dict[str, Any] | None = None,
    ) -> DocumentSection:
        return DocumentSection(
            section_id=stable_id(
                document_id,
                "section",
                str(index),
            ),
            document_id=document_id,
            page_number=page_number,
            section_title=title,
            text=normalize_text(text),
            metadata=metadata or {},
        )

    # --------------------------------------------------------
    # PDF
    # --------------------------------------------------------

    def _parse_pdf(
        self,
        path: Path,
        document_id: str,
    ) -> list[DocumentSection]:
        from pypdf import PdfReader

        reader = PdfReader(
            str(path),
            strict=False,
        )

        if reader.is_encrypted:
            raise ValueError(
                "Encrypted PDFs are not supported"
            )

        sections = []

        for index, page in enumerate(
            reader.pages,
            start=1,
        ):
            text = page.extract_text() or ""

            if (
                not text.strip()
                and self.config.enable_ocr
            ):
                text = self._ocr_pdf_page(
                    path,
                    index,
                )

            text = normalize_text(text)

            if not text:
                continue

            sections.append(
                self._section(
                    document_id=document_id,
                    index=index,
                    text=text,
                    page_number=index,
                    title=f"Page {index}",
                )
            )

        return sections

    def _ocr_pdf_page(
        self,
        path: Path,
        page_number: int,
    ) -> str:
        try:
            import pytesseract

            from pdf2image import (
                convert_from_path,
            )
        except ImportError as exc:
            raise RuntimeError(
                "OCR requires pytesseract "
                "and pdf2image"
            ) from exc

        images = convert_from_path(
            str(path),
            first_page=page_number,
            last_page=page_number,
            dpi=200,
        )

        if not images:
            return ""

        return pytesseract.image_to_string(
            images[0]
        )

    # --------------------------------------------------------
    # DOCX
    # --------------------------------------------------------

    def _parse_docx(
        self,
        path: Path,
        document_id: str,
    ) -> list[DocumentSection]:
        from docx import Document

        document = Document(
            str(path)
        )

        blocks = []

        for paragraph in document.paragraphs:
            if paragraph.text.strip():
                blocks.append(
                    paragraph.text
                )

        for table_index, table in enumerate(
            document.tables,
            start=1,
        ):
            blocks.append(
                f"Table {table_index}"
            )

            for row in table.rows:
                cells = [
                    normalize_text(cell.text)
                    for cell in row.cells
                ]

                blocks.append(
                    " | ".join(cells)
                )

        return [
            self._section(
                document_id,
                1,
                "\n".join(blocks),
                title=path.stem,
            )
        ]

    # --------------------------------------------------------
    # TEXT / MARKDOWN
    # --------------------------------------------------------

    def _parse_text(
        self,
        path: Path,
        document_id: str,
    ) -> list[DocumentSection]:
        text = path.read_text(
            encoding="utf-8",
            errors="replace",
        )

        return [
            self._section(
                document_id,
                1,
                text,
                title=path.stem,
            )
        ]

    # --------------------------------------------------------
    # HTML
    # --------------------------------------------------------

    def _parse_html(
        self,
        path: Path,
        document_id: str,
    ) -> list[DocumentSection]:
        from bs4 import BeautifulSoup

        html = path.read_text(
            encoding="utf-8",
            errors="replace",
        )

        soup = BeautifulSoup(
            html,
            "html.parser",
        )

        for tag in soup([
            "script",
            "style",
            "noscript",
            "iframe",
        ]):
            tag.decompose()

        text = soup.get_text(
            separator="\n",
        )

        return [
            self._section(
                document_id,
                1,
                text,
                title=path.stem,
            )
        ]

    # --------------------------------------------------------
    # CSV
    # --------------------------------------------------------

    def _parse_csv(
        self,
        path: Path,
        document_id: str,
    ) -> list[DocumentSection]:
        import pandas as pd

        dataframe = pd.read_csv(
            path,
            dtype=str,
            keep_default_na=False,
        )

        lines = []

        for index, row in dataframe.iterrows():
            values = [
                f"{column}: {row[column]}"
                for column in dataframe.columns
            ]

            lines.append(
                f"Row {index + 1}: "
                + " | ".join(values)
            )

        return [
            self._section(
                document_id,
                1,
                "\n".join(lines),
                title=path.stem,
            )
        ]

    # --------------------------------------------------------
    # XLSX
    # --------------------------------------------------------

    def _parse_xlsx(
        self,
        path: Path,
        document_id: str,
    ) -> list[DocumentSection]:
        from openpyxl import load_workbook

        workbook = load_workbook(
            path,
            read_only=True,
            data_only=True,
        )

        sections = []

        try:
            for index, worksheet in enumerate(
                workbook.worksheets,
                start=1,
            ):
                lines = []

                for row_number, row in enumerate(
                    worksheet.iter_rows(
                        values_only=True
                    ),
                    start=1,
                ):
                    values = [
                        str(value)
                        if value is not None
                        else ""
                        for value in row
                    ]

                    if any(values):
                        lines.append(
                            f"Row {row_number}: "
                            + " | ".join(values)
                        )

                sections.append(
                    self._section(
                        document_id,
                        index,
                        "\n".join(lines),
                        title=worksheet.title,
                        metadata={
                            "sheet_name": (
                                worksheet.title
                            )
                        },
                    )
                )
        finally:
            workbook.close()

        return sections

    # --------------------------------------------------------
    # JSON
    # --------------------------------------------------------

    def _parse_json(
        self,
        path: Path,
        document_id: str,
    ) -> list[DocumentSection]:
        data = json.loads(
            path.read_text(
                encoding="utf-8"
            )
        )

        text = json.dumps(
            data,
            ensure_ascii=False,
            indent=2,
        )

        return [
            self._section(
                document_id,
                1,
                text,
                title=path.stem,
            )
        ]


# ============================================================
# TOKEN-AWARE CHUNKER
# ============================================================

class DocumentChunker:
    """
    Uses the embedding model tokenizer when available.

    Without transformers, falls back to a simple
    whitespace token approximation.

    Chunk sizes in fallback mode are word counts,
    not true model token counts.
    """

    def __init__(
        self,
        config: IngestionConfig,
    ) -> None:
        self.config = config

        self.tokenizer = None

        if config.enable_embeddings:
            try:
                from transformers import (
                    AutoTokenizer,
                )

                self.tokenizer = (
                    AutoTokenizer.from_pretrained(
                        config.embedding_model
                    )
                )
            except Exception as exc:
                LOGGER.warning(
                    "Tokenizer unavailable; "
                    "using word-based chunking: %s",
                    exc,
                )

    def _encode(
        self,
        text: str,
    ) -> list[Any]:
        if self.tokenizer is not None:
            return self.tokenizer.encode(
                text,
                add_special_tokens=False,
            )

        return text.split()

    def _decode(
        self,
        tokens: list[Any],
    ) -> str:
        if self.tokenizer is not None:
            return self.tokenizer.decode(
                tokens,
                skip_special_tokens=True,
            )

        return " ".join(tokens)

    def chunk(
        self,
        section: DocumentSection,
        document: DocumentMetadata,
        start_index: int = 0,
    ) -> list[DocumentChunk]:
        tokens = self._encode(
            section.text
        )

        if not tokens:
            return []

        size = self.config.chunk_size
        overlap = self.config.chunk_overlap

        if size <= 0:
            raise ValueError(
                "chunk_size must be positive"
            )

        if overlap < 0 or overlap >= size:
            raise ValueError(
                "chunk_overlap must be >= 0 "
                "and less than chunk_size"
            )

        chunks = []

        start = 0
        index = start_index

        while start < len(tokens):
            end = min(
                start + size,
                len(tokens),
            )

            window = tokens[start:end]

            chunk_text = normalize_text(
                self._decode(window)
            )

            if chunk_text:
                chunk_id = stable_id(
                    document.document_id,
                    section.section_id,
                    str(start),
                    str(end),
                )

                chunks.append(
                    DocumentChunk(
                        chunk_id=chunk_id,
                        document_id=(
                            document.document_id
                        ),
                        tenant_id=(
                            document.tenant_id
                        ),
                        chunk_index=index,
                        text=chunk_text,
                        token_count=len(window),
                        page_number=(
                            section.page_number
                        ),
                        source_path=(
                            document.source_path
                        ),
                        document_type=(
                            document.document_type
                        ),
                        equipment_type=(
                            document.equipment_type
                        ),
                        metadata={
                            "filename": (
                                document.filename
                            ),
                            "title": (
                                document.title
                            ),
                            "section_title": (
                                section.section_title
                            ),
                            "sha256": (
                                document.sha256
                            ),
                            "token_start": start,
                            "token_end": end,
                            **section.metadata,
                        },
                    )
                )

                index += 1

            if end == len(tokens):
                break

            start = end - overlap

        return chunks


# ============================================================
# EMBEDDING SERVICE
# ============================================================

class EmbeddingService:
    def __init__(
        self,
        model_name: str,
    ) -> None:
        from sentence_transformers import (
            SentenceTransformer,
        )

        self.model = SentenceTransformer(
            model_name
        )

    def encode(
        self,
        texts: list[str],
        batch_size: int = 32,
    ) -> list[list[float]]:
        if not texts:
            return []

        vectors = self.model.encode(
            texts,
            batch_size=batch_size,
            normalize_embeddings=True,
            show_progress_bar=False,
        )

        return vectors.tolist()

    def dimension(self) -> int:
        dimension = (
            self.model.get_sentence_embedding_dimension()
        )

        if dimension is None:
            raise RuntimeError(
                "Unknown embedding dimension"
            )

        return int(dimension)


# ============================================================
# QDRANT VECTOR STORE
# ============================================================

class QdrantDocumentStore:
    def __init__(
        self,
        url: str,
        collection_name: str,
        vector_size: int,
    ) -> None:
        from qdrant_client import (
            QdrantClient,
            models,
        )

        self.models = models

        self.client = QdrantClient(
            url=url
        )

        self.collection_name = collection_name

        if not self.client.collection_exists(
            collection_name
        ):
            self.client.create_collection(
                collection_name=collection_name,
                vectors_config=(
                    models.VectorParams(
                        size=vector_size,
                        distance=(
                            models.Distance.COSINE
                        ),
                    )
                ),
            )
        else:
            collection = (
                self.client.get_collection(
                    collection_name
                )
            )

            existing_vectors = (
                collection.config.params.vectors
            )

            if (
                hasattr(existing_vectors, "size")
                and existing_vectors.size
                != vector_size
            ):
                raise ValueError(
                    "Existing Qdrant collection "
                    "uses a different vector size"
                )

    def upsert(
        self,
        chunks: list[DocumentChunk],
        vectors: list[list[float]],
        batch_size: int = 32,
    ) -> None:
        if len(chunks) != len(vectors):
            raise ValueError(
                "Chunks and vectors length mismatch"
            )

        for start in range(
            0,
            len(chunks),
            batch_size,
        ):
            chunk_batch = chunks[
                start:start + batch_size
            ]

            vector_batch = vectors[
                start:start + batch_size
            ]

            points = []

            for chunk, vector in zip(
                chunk_batch,
                vector_batch,
            ):
                payload = chunk.model_dump()

                points.append(
                    self.models.PointStruct(
                        id=chunk.chunk_id,
                        vector=vector,
                        payload=payload,
                    )
                )

            self.client.upsert(
                collection_name=(
                    self.collection_name
                ),
                points=points,
                wait=True,
            )


# ============================================================
# JSONL WRITER
# ============================================================

class JSONLWriter:
    @staticmethod
    def write(
        path: Path,
        records: Iterable[
            BaseModel | dict[str, Any]
        ],
    ) -> None:
        path.parent.mkdir(
            parents=True,
            exist_ok=True,
        )

        with path.open(
            "w",
            encoding="utf-8",
        ) as stream:
            for record in records:
                if isinstance(
                    record,
                    BaseModel,
                ):
                    data = record.model_dump()
                else:
                    data = record

                stream.write(
                    json.dumps(
                        data,
                        ensure_ascii=False,
                        default=str,
                    )
                    + "\n"
                )


# ============================================================
# INGESTION PIPELINE
# ============================================================

class IndustrialDocumentIngestionPipeline:
    def __init__(
        self,
        config: IngestionConfig,
    ) -> None:
        self.config = config

        self.discovery = DocumentDiscovery(
            config
        )

        self.parser = DocumentParser(
            config
        )

        self.chunker = DocumentChunker(
            config
        )

        self.embedding_service = None
        self.vector_store = None

        if config.enable_qdrant and not (
            config.enable_embeddings
        ):
            raise ValueError(
                "Qdrant indexing requires embeddings"
            )

        if config.enable_embeddings:
            self.embedding_service = (
                EmbeddingService(
                    config.embedding_model
                )
            )

        if config.enable_qdrant:
            self.vector_store = (
                QdrantDocumentStore(
                    url=config.qdrant_url,
                    collection_name=(
                        config.qdrant_collection
                    ),
                    vector_size=(
                        self.embedding_service.dimension()
                    ),
                )
            )

    # --------------------------------------------------------
    # DOCUMENT METADATA
    # --------------------------------------------------------

    def build_metadata(
        self,
        path: Path,
    ) -> DocumentMetadata:
        digest = sha256_file(
            path
        )

        document_id = stable_id(
            self.config.tenant_id,
            digest,
        )

        return DocumentMetadata(
            document_id=document_id,
            tenant_id=self.config.tenant_id,
            filename=path.name,
            source_path=str(
                path.resolve()
            ),
            extension=path.suffix.lower(),
            file_size_bytes=(
                path.stat().st_size
            ),
            sha256=digest,
            title=path.stem,
            equipment_type=(
                infer_equipment_type(
                    path.name
                )
            ),
            document_type=(
                infer_document_type(
                    path.name
                )
            ),
            created_at=utc_now(),
        )

    # --------------------------------------------------------
    # INGEST ONE DOCUMENT
    # --------------------------------------------------------

    def ingest_document(
        self,
        path: Path,
    ) -> tuple[
        DocumentMetadata,
        list[DocumentSection],
        list[DocumentChunk],
    ]:
        if path.is_symlink():
            raise ValueError(
                "Symbolic links are not allowed"
            )

        if (
            path.suffix.lower()
            not in self.config.allowed_extensions
        ):
            raise ValueError(
                "Unsupported document extension"
            )

        if (
            path.stat().st_size
            > self.config.max_file_size_mb
            * 1024
            * 1024
        ):
            raise ValueError(
                "File exceeds configured size limit"
            )

        metadata = self.build_metadata(
            path
        )

        sections = self.parser.parse(
            path,
            metadata.document_id,
        )

        total_characters = sum(
            len(section.text)
            for section in sections
        )

        if (
            total_characters
            > self.config.max_extracted_characters
        ):
            raise ValueError(
                "Extracted text exceeds size limit"
            )

        chunks = []

        for section in sections:
            new_chunks = self.chunker.chunk(
                section,
                metadata,
                start_index=len(chunks),
            )

            chunks.extend(
                new_chunks
            )

        return (
            metadata,
            sections,
            chunks,
        )

    # --------------------------------------------------------
    # INGEST ALL DOCUMENTS
    # --------------------------------------------------------

    def run(self) -> dict[str, Any]:
        paths = self.discovery.discover()

        output_dir = Path(
            self.config.output_dir
        )

        output_dir.mkdir(
            parents=True,
            exist_ok=True,
        )

        documents = []
        sections = []
        chunks = []
        results = []

        seen_hashes = set()

        LOGGER.info(
            "Discovered %d files",
            len(paths),
        )

        for path in paths:
            try:
                digest = sha256_file(
                    path
                )

                if digest in seen_hashes:
                    LOGGER.info(
                        "Skipping duplicate: %s",
                        path,
                    )
                    continue

                (
                    metadata,
                    document_sections,
                    document_chunks,
                ) = self.ingest_document(
                    path
                )

                seen_hashes.add(
                    digest
                )

                documents.append(
                    metadata
                )

                sections.extend(
                    document_sections
                )

                chunks.extend(
                    document_chunks
                )

                results.append(
                    IngestionResult(
                        document_id=(
                            metadata.document_id
                        ),
                        filename=path.name,
                        status="success",
                        section_count=len(
                            document_sections
                        ),
                        chunk_count=len(
                            document_chunks
                        ),
                    )
                )

                LOGGER.info(
                    "Ingested %s: %d chunks",
                    path.name,
                    len(document_chunks),
                )

            except Exception as exc:
                LOGGER.exception(
                    "Failed to ingest %s",
                    path,
                )

                results.append(
                    IngestionResult(
                        document_id="",
                        filename=path.name,
                        status="failed",
                        error=str(exc),
                    )
                )

        # ----------------------------------------------------
        # EXPORT DOCUMENT RECORDS
        # ----------------------------------------------------

        JSONLWriter.write(
            output_dir / "documents.jsonl",
            documents,
        )

        JSONLWriter.write(
            output_dir / "sections.jsonl",
            sections,
        )

        JSONLWriter.write(
            output_dir / "chunks.jsonl",
            chunks,
        )

        JSONLWriter.write(
            output_dir / "ingestion_results.jsonl",
            results,
        )

        # ----------------------------------------------------
        # EMBEDDINGS
        # ----------------------------------------------------

        embedding_count = 0

        if (
            self.embedding_service is not None
            and chunks
        ):
            LOGGER.info(
                "Generating embeddings for %d chunks",
                len(chunks),
            )

            vectors = (
                self.embedding_service.encode(
                    [chunk.text for chunk in chunks],
                    batch_size=(
                        self.config.batch_size
                    ),
                )
            )

            embedding_count = len(
                vectors
            )

            # Embeddings are indexed in Qdrant when
            # enabled. Otherwise export as JSONL.
            if self.vector_store is not None:
                self.vector_store.upsert(
                    chunks,
                    vectors,
                    batch_size=(
                        self.config.batch_size
                    ),
                )
            else:
                JSONLWriter.write(
                    output_dir / "embeddings.jsonl",
                    (
                        {
                            "chunk_id": chunk.chunk_id,
                            "vector": vector,
                            "model": (
                                self.config.embedding_model
                            ),
                        }
                        for chunk, vector in zip(
                            chunks,
                            vectors,
                        )
                    ),
                )

        # ----------------------------------------------------
        # REPORT
        # ----------------------------------------------------

        report = {
            "generated_at": utc_now(),
            "tenant_id": (
                self.config.tenant_id
            ),
            "files_discovered": len(paths),
            "documents_ingested": len(
                documents
            ),
            "documents_failed": sum(
                result.status == "failed"
                for result in results
            ),
            "sections_created": len(
                sections
            ),
            "chunks_created": len(
                chunks
            ),
            "embeddings_created": (
                embedding_count
            ),
            "qdrant_indexing_enabled": (
                self.config.enable_qdrant
            ),
            "qdrant_collection": (
                self.config.qdrant_collection
                if self.config.enable_qdrant
                else None
            ),
            "configuration": asdict(
                self.config
            ),
        }

        report_path = (
            output_dir / "ingestion_report.json"
        )

        report_path.write_text(
            json.dumps(
                report,
                indent=2,
                ensure_ascii=False,
            ),
            encoding="utf-8",
        )

        LOGGER.info(
            "Ingestion complete: %s",
            report_path,
        )

        return report


# ============================================================
# CLI
# ============================================================

def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Ingest industrial documents "
            "for RAG retrieval"
        )
    )

    parser.add_argument(
        "--input-dir",
        default="data/documents",
    )

    parser.add_argument(
        "--output-dir",
        default="data/processed/documents",
    )

    parser.add_argument(
        "--tenant-id",
        default="default",
    )

    parser.add_argument(
        "--chunk-size",
        type=int,
        default=450,
    )

    parser.add_argument(
        "--chunk-overlap",
        type=int,
        default=75,
    )

    parser.add_argument(
        "--max-file-size-mb",
        type=int,
        default=25,
    )

    parser.add_argument(
        "--ocr",
        action="store_true",
    )

    parser.add_argument(
        "--embeddings",
        action="store_true",
    )

    parser.add_argument(
        "--qdrant",
        action="store_true",
    )

    parser.add_argument(
        "--embedding-model",
        default=(
            "sentence-transformers/"
            "all-MiniLM-L6-v2"
        ),
    )

    parser.add_argument(
        "--qdrant-url",
        default="http://localhost:6333",
    )

    parser.add_argument(
        "--qdrant-collection",
        default="industrial_documents",
    )

    parser.add_argument(
        "--batch-size",
        type=int,
        default=32,
    )

    parser.add_argument(
        "--no-recursive",
        action="store_true",
    )

    parser.add_argument(
        "--verbose",
        action="store_true",
    )

    return parser.parse_args()


def main() -> None:
    args = parse_args()

    logging.basicConfig(
        level=(
            logging.DEBUG
            if args.verbose
            else logging.INFO
        ),
        format=(
            "%(asctime)s | %(levelname)s | "
            "%(name)s | %(message)s"
        ),
    )

    config = IngestionConfig(
        input_dir=args.input_dir,
        output_dir=args.output_dir,
        tenant_id=args.tenant_id,
        chunk_size=args.chunk_size,
        chunk_overlap=args.chunk_overlap,
        enable_ocr=args.ocr,
        enable_embeddings=(
            args.embeddings or args.qdrant
        ),
        enable_qdrant=args.qdrant,
        embedding_model=args.embedding_model,
        qdrant_url=args.qdrant_url,
        qdrant_collection=(
            args.qdrant_collection
        ),
        batch_size=args.batch_size,
        recursive=not args.no_recursive,
    )

    pipeline = (
        IndustrialDocumentIngestionPipeline(
            config
        )
    )

    report = pipeline.run()

    print(
        json.dumps(
            report,
            indent=2,
        )
    )


if __name__ == "__main__":
    main()