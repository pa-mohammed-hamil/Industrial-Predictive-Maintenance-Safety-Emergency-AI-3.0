
"""
app/ai/rag/ingestion/document_processing.py

Industrial Document Processing Engine
=====================================

Python 3.11+

Dependencies:
    pip install pypdf python-docx beautifulsoup4

Optional OCR:
    pip install pytesseract pillow pdf2image

OCR system dependencies:
    Tesseract OCR
    Poppler (for scanned PDF rendering)

Features:
- PDF text extraction
- DOCX paragraph and table extraction
- TXT and Markdown processing
- CSV table extraction
- JSON processing
- HTML text extraction
- PNG/JPG/TIFF OCR
- Scanned PDF OCR fallback
- Page-level metadata
- Document hashing
- Text normalization
- Industrial entity extraction
- Document classification
- Batch document processing
- File size limits
- Safe path handling
- Citation-ready page records

Security:
- Do not trust uploaded filenames or MIME types.
- Run document parsers in an isolated worker.
- Scan uploaded files before processing.
- Enforce tenant and document permissions.
- Do not execute embedded scripts or macros.
"""

from __future__ import annotations

import asyncio
import csv
import hashlib
import html
import io
import json
import logging
import mimetypes
import re

from dataclasses import dataclass, field
from datetime import datetime, timezone
from enum import Enum
from pathlib import Path
from typing import Any, Mapping, Sequence


logger = logging.getLogger(__name__)


class DocumentProcessingError(Exception):
    """Base document processing exception."""


class UnsupportedDocumentError(DocumentProcessingError):
    """Unsupported document format."""


class DocumentValidationError(DocumentProcessingError):
    """Invalid document input."""


class DocumentExtractionError(DocumentProcessingError):
    """Document text extraction failed."""


class DocumentType(str, Enum):
    MAINTENANCE_MANUAL = "maintenance_manual"
    SAFETY_PROCEDURE = "safety_procedure"
    INSPECTION_REPORT = "inspection_report"
    INCIDENT_REPORT = "incident_report"
    TROUBLESHOOTING = "troubleshooting"
    TECHNICAL_SPECIFICATION = "technical_specification"
    EQUIPMENT_MANUAL = "equipment_manual"
    WORK_ORDER = "work_order"
    GENERAL = "general"


@dataclass(slots=True)
class DocumentPage:
    page_number: int
    text: str
    metadata: dict[str, Any] = field(default_factory=dict)

    def to_dict(self) -> dict[str, Any]:
        return {
            "page_number": self.page_number,
            "text": self.text,
            "metadata": dict(self.metadata),
        }


@dataclass(slots=True)
class ProcessedDocument:
    document_id: str
    company_id: str
    filename: str
    file_type: str
    document_type: str
    content_hash: str
    text: str
    pages: list[DocumentPage]
    metadata: dict[str, Any]
    entities: dict[str, list[str]]
    processed_at: datetime

    def to_dict(self) -> dict[str, Any]:
        return {
            "document_id": self.document_id,
            "company_id": self.company_id,
            "filename": self.filename,
            "file_type": self.file_type,
            "document_type": self.document_type,
            "content_hash": self.content_hash,
            "text": self.text,
            "pages": [
                page.to_dict()
                for page in self.pages
            ],
            "metadata": dict(self.metadata),
            "entities": dict(self.entities),
            "processed_at": self.processed_at.isoformat(),
        }


@dataclass(slots=True)
class DocumentProcessingConfig:
    max_file_size_mb: int = 30
    max_pages: int = 300
    max_text_chars: int = 3_000_000
    enable_ocr: bool = False
    ocr_language: str = "eng"
    normalize_whitespace: bool = True
    extract_entities: bool = True
    classify_documents: bool = True

    def __post_init__(self) -> None:
        for name in (
            "max_file_size_mb",
            "max_pages",
            "max_text_chars",
        ):
            value = getattr(self, name)

            if (
                isinstance(value, bool)
                or not isinstance(value, int)
                or value <= 0
            ):
                raise DocumentValidationError(
                    f"{name} must be a positive integer"
                )


class IndustrialEntityExtractor:
    """
    Rule-based industrial entity extraction.

    This is a baseline extractor, not a trained NER model.
    """

    PATTERNS = {
        "asset_ids": re.compile(
            r"\b(?:MOTOR|PUMP|VALVE|COMP|FAN|"
            r"ASSET|MACHINE)[-_ ]?\d{2,8}\b",
            re.IGNORECASE,
        ),
        "temperatures": re.compile(
            r"\b\d+(?:\.\d+)?\s?(?:°C|°F|Celsius|Fahrenheit)\b",
            re.IGNORECASE,
        ),
        "pressures": re.compile(
            r"\b\d+(?:\.\d+)?\s?(?:bar|psi|kPa|MPa)\b",
            re.IGNORECASE,
        ),
        "vibration": re.compile(
            r"\b\d+(?:\.\d+)?\s?(?:mm/s|m/s²|g)\b",
            re.IGNORECASE,
        ),
        "voltage": re.compile(
            r"\b\d+(?:\.\d+)?\s?(?:kV|V)\b"
        ),
        "current": re.compile(
            r"\b\d+(?:\.\d+)?\s?(?:mA|A)\b"
        ),
        "power": re.compile(
            r"\b\d+(?:\.\d+)?\s?(?:kW|MW|W)\b"
        ),
        "fault_codes": re.compile(
            r"\b(?:FAULT|ERROR|DTC)[-_ :]*[A-Z0-9-]{2,12}\b",
            re.IGNORECASE,
        ),
    }

    def extract(self, text: str) -> dict[str, list[str]]:
        results: dict[str, list[str]] = {}

        for name, pattern in self.PATTERNS.items():
            matches = pattern.findall(text)

            results[name] = sorted(
                set(matches),
                key=str.casefold,
            )

        return results


class IndustrialDocumentClassifier:
    """
    Lightweight keyword-based document classifier.

    Replace with a trained classifier when needed.
    """

    KEYWORDS = {
        DocumentType.MAINTENANCE_MANUAL: (
            "preventive maintenance",
            "lubrication schedule",
            "maintenance interval",
            "bearing replacement",
        ),
        DocumentType.SAFETY_PROCEDURE: (
            "safety procedure",
            "lockout tagout",
            "hazard assessment",
            "personal protective equipment",
        ),
        DocumentType.INSPECTION_REPORT: (
            "inspection report",
            "inspection findings",
            "visual inspection",
            "defect severity",
        ),
        DocumentType.INCIDENT_REPORT: (
            "incident report",
            "root cause analysis",
            "near miss",
            "corrective action",
        ),
        DocumentType.TROUBLESHOOTING: (
            "troubleshooting",
            "fault diagnosis",
            "possible causes",
            "symptom",
        ),
        DocumentType.TECHNICAL_SPECIFICATION: (
            "technical specification",
            "rated voltage",
            "rated power",
            "operating limits",
        ),
        DocumentType.EQUIPMENT_MANUAL: (
            "equipment manual",
            "installation guide",
            "operating instructions",
        ),
        DocumentType.WORK_ORDER: (
            "work order",
            "assigned technician",
            "job completion",
        ),
    }

    def classify(self, text: str) -> str:
        sample = text[:30000].lower()

        scores = {
            doc_type: sum(
                sample.count(keyword)
                for keyword in keywords
            )
            for doc_type, keywords
            in self.KEYWORDS.items()
        }

        best = max(
            scores,
            key=scores.get,
        )

        if scores[best] == 0:
            return DocumentType.GENERAL.value

        return best.value


class DocumentProcessor:
    SUPPORTED_EXTENSIONS = frozenset({
        ".pdf",
        ".docx",
        ".txt",
        ".md",
        ".csv",
        ".json",
        ".html",
        ".htm",
        ".png",
        ".jpg",
        ".jpeg",
        ".tif",
        ".tiff",
    })

    def __init__(
        self,
        config: DocumentProcessingConfig | None = None,
    ) -> None:
        self.config = config or DocumentProcessingConfig()
        self.entity_extractor = IndustrialEntityExtractor()
        self.classifier = IndustrialDocumentClassifier()

    # =====================================================
    # HELPERS
    # =====================================================

    @staticmethod
    def _require_text(value: str, name: str) -> str:
        if not isinstance(value, str) or not value.strip():
            raise DocumentValidationError(
                f"{name} must be a nonempty string"
            )
        return value.strip()

    @staticmethod
    def _hash_bytes(data: bytes) -> str:
        return hashlib.sha256(data).hexdigest()

    @staticmethod
    def _normalize(text: str) -> str:
        text = text.replace("\r\n", "\n")
        text = text.replace("\r", "\n")
        text = text.replace("\x00", "")

        text = re.sub(
            r"[ \t]+",
            " ",
            text,
        )

        text = re.sub(
            r"\n{3,}",
            "\n\n",
            text,
        )

        return text.strip()

    def _validate_size(self, data: bytes) -> None:
        if not data:
            raise DocumentValidationError(
                "Document is empty"
            )

        maximum = (
            self.config.max_file_size_mb
            * 1024
            * 1024
        )

        if len(data) > maximum:
            raise DocumentValidationError(
                "Document exceeds maximum file size"
            )

    def _validate_pages(
        self,
        pages: Sequence[DocumentPage],
    ) -> None:
        if len(pages) > self.config.max_pages:
            raise DocumentValidationError(
                "Document exceeds maximum page count"
            )

        total_chars = sum(
            len(page.text)
            for page in pages
        )

        if total_chars > self.config.max_text_chars:
            raise DocumentValidationError(
                "Extracted text exceeds maximum length"
            )

    # =====================================================
    # TEXT
    # =====================================================

    @staticmethod
    def _decode_text(data: bytes) -> str:
        try:
            return data.decode("utf-8-sig")
        except UnicodeDecodeError:
            raise DocumentExtractionError(
                "Text file must use UTF-8 encoding"
            ) from None

    def _process_text(
        self,
        data: bytes,
    ) -> list[DocumentPage]:
        return [
            DocumentPage(
                page_number=1,
                text=self._decode_text(data),
            )
        ]

    # =====================================================
    # PDF
    # =====================================================

    def _process_pdf(
        self,
        data: bytes,
    ) -> list[DocumentPage]:
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise DocumentExtractionError(
                "Install pypdf to process PDFs"
            ) from exc

        try:
            reader = PdfReader(io.BytesIO(data))

            if reader.is_encrypted:
                raise DocumentExtractionError(
                    "Encrypted PDFs are not supported"
                )

            if len(reader.pages) > self.config.max_pages:
                raise DocumentValidationError(
                    "PDF exceeds maximum page count"
                )

            pages: list[DocumentPage] = []

            for index, page in enumerate(
                reader.pages,
                start=1,
            ):
                text = page.extract_text() or ""
                extraction_method = "pdf_text"

                if (
                    not text.strip()
                    and self.config.enable_ocr
                ):
                    text = self._ocr_pdf_page(
                        data,
                        index,
                    )
                    extraction_method = "ocr"

                pages.append(
                    DocumentPage(
                        page_number=index,
                        text=text,
                        metadata={
                            "extraction_method":
                                extraction_method,
                        },
                    )
                )

            return pages

        except DocumentProcessingError:
            raise
        except Exception as exc:
            raise DocumentExtractionError(
                "PDF extraction failed"
            ) from exc

    # =====================================================
    # PDF OCR
    # =====================================================

    def _ocr_pdf_page(
        self,
        data: bytes,
        page_number: int,
    ) -> str:
        try:
            from pdf2image import convert_from_bytes
            import pytesseract
        except ImportError as exc:
            raise DocumentExtractionError(
                "Install pdf2image and pytesseract"
            ) from exc

        images = convert_from_bytes(
            data,
            first_page=page_number,
            last_page=page_number,
            dpi=200,
            thread_count=1,
        )

        if not images:
            return ""

        try:
            return pytesseract.image_to_string(
                images[0],
                lang=self.config.ocr_language,
            )
        finally:
            for image in images:
                image.close()

    # =====================================================
    # DOCX
    # =====================================================

    def _process_docx(
        self,
        data: bytes,
    ) -> list[DocumentPage]:
        try:
            from docx import Document
        except ImportError as exc:
            raise DocumentExtractionError(
                "Install python-docx"
            ) from exc

        try:
            document = Document(io.BytesIO(data))

            parts = []

            for paragraph in document.paragraphs:
                if paragraph.text.strip():
                    parts.append(
                        paragraph.text
                    )

            for table in document.tables:
                for row in table.rows:
                    values = [
                        cell.text.strip()
                        for cell in row.cells
                    ]

                    parts.append(
                        " | ".join(values)
                    )

            return [
                DocumentPage(
                    page_number=1,
                    text="\n".join(parts),
                    metadata={
                        "page_number_is_logical": True,
                        "extraction_method": "docx",
                    },
                )
            ]

        except Exception as exc:
            raise DocumentExtractionError(
                "DOCX extraction failed"
            ) from exc

    # =====================================================
    # CSV
    # =====================================================

    def _process_csv(
        self,
        data: bytes,
    ) -> list[DocumentPage]:
        text = self._decode_text(data)

        reader = csv.reader(
            io.StringIO(text)
        )

        lines = []

        for row in reader:
            lines.append(
                " | ".join(
                    str(value)
                    for value in row
                )
            )

        return [
            DocumentPage(
                page_number=1,
                text="\n".join(lines),
                metadata={
                    "extraction_method": "csv",
                },
            )
        ]

    # =====================================================
    # JSON
    # =====================================================

    def _process_json(
        self,
        data: bytes,
    ) -> list[DocumentPage]:
        try:
            payload = json.loads(
                self._decode_text(data)
            )
        except json.JSONDecodeError as exc:
            raise DocumentExtractionError(
                "Invalid JSON document"
            ) from exc

        return [
            DocumentPage(
                page_number=1,
                text=json.dumps(
                    payload,
                    ensure_ascii=False,
                    indent=2,
                ),
                metadata={
                    "extraction_method": "json",
                },
            )
        ]

    # =====================================================
    # HTML
    # =====================================================

    def _process_html(
        self,
        data: bytes,
    ) -> list[DocumentPage]:
        try:
            from bs4 import BeautifulSoup
        except ImportError as exc:
            raise DocumentExtractionError(
                "Install beautifulsoup4"
            ) from exc

        soup = BeautifulSoup(
            self._decode_text(data),
            "html.parser",
        )

        for element in soup([
            "script",
            "style",
            "noscript",
            "iframe",
        ]):
            element.decompose()

        text = soup.get_text(
            separator="\n",
            strip=True,
        )

        return [
            DocumentPage(
                page_number=1,
                text=html.unescape(text),
                metadata={
                    "extraction_method": "html",
                },
            )
        ]

    # =====================================================
    # IMAGE OCR
    # =====================================================

    def _process_image(
        self,
        data: bytes,
    ) -> list[DocumentPage]:
        if not self.config.enable_ocr:
            raise DocumentExtractionError(
                "Image processing requires OCR enabled"
            )

        try:
            from PIL import Image
            import pytesseract
        except ImportError as exc:
            raise DocumentExtractionError(
                "Install pillow and pytesseract"
            ) from exc

        try:
            with Image.open(
                io.BytesIO(data)
            ) as image:
                image.load()

                text = pytesseract.image_to_string(
                    image,
                    lang=self.config.ocr_language,
                )

            return [
                DocumentPage(
                    page_number=1,
                    text=text,
                    metadata={
                        "extraction_method": "ocr",
                    },
                )
            ]

        except Exception as exc:
            raise DocumentExtractionError(
                "Image OCR failed"
            ) from exc

    # =====================================================
    # EXTRACTION ROUTER
    # =====================================================

    def _extract(
        self,
        data: bytes,
        extension: str,
    ) -> list[DocumentPage]:
        if extension == ".pdf":
            return self._process_pdf(data)

        if extension == ".docx":
            return self._process_docx(data)

        if extension in {".txt", ".md"}:
            return self._process_text(data)

        if extension == ".csv":
            return self._process_csv(data)

        if extension == ".json":
            return self._process_json(data)

        if extension in {".html", ".htm"}:
            return self._process_html(data)

        if extension in {
            ".png",
            ".jpg",
            ".jpeg",
            ".tif",
            ".tiff",
        }:
            return self._process_image(data)

        raise UnsupportedDocumentError(
            f"Unsupported extension: {extension}"
        )

    # =====================================================
    # PROCESS BYTES
    # =====================================================

    def process_bytes(
        self,
        *,
        data: bytes,
        filename: str,
        company_id: str,
        document_id: str | None = None,
        metadata: Mapping[str, Any] | None = None,
    ) -> ProcessedDocument:
        filename = self._require_text(
            filename,
            "filename",
        )

        company_id = self._require_text(
            company_id,
            "company_id",
        )

        if not isinstance(data, bytes):
            raise DocumentValidationError(
                "data must be bytes"
            )

        self._validate_size(data)

        safe_filename = Path(
            filename.replace("\\", "/")
        ).name

        extension = Path(
            safe_filename
        ).suffix.lower()

        if extension not in self.SUPPORTED_EXTENSIONS:
            raise UnsupportedDocumentError(
                f"Unsupported document type: {extension}"
            )

        content_hash = self._hash_bytes(data)

        if document_id is None:
            document_id = (
                f"doc_{content_hash[:24]}"
            )

        pages = self._extract(
            data,
            extension,
        )

        self._validate_pages(pages)

        if self.config.normalize_whitespace:
            for page in pages:
                page.text = self._normalize(
                    page.text
                )

        full_text = "\n\n".join(
            page.text
            for page in pages
            if page.text.strip()
        )

        if len(full_text) > self.config.max_text_chars:
            raise DocumentValidationError(
                "Document exceeds maximum text length"
            )

        document_type = (
            self.classifier.classify(full_text)
            if self.config.classify_documents
            else DocumentType.GENERAL.value
        )

        entities = (
            self.entity_extractor.extract(full_text)
            if self.config.extract_entities
            else {}
        )

        document_metadata = dict(
            metadata or {}
        )

        document_metadata.update({
            "filename": safe_filename,
            "file_extension": extension,
            "mime_type": mimetypes.guess_type(
                safe_filename
            )[0],
            "page_count": len(pages),
            "character_count": len(full_text),
            "content_hash": content_hash,
            "document_type": document_type,
        })

        return ProcessedDocument(
            document_id=document_id,
            company_id=company_id,
            filename=safe_filename,
            file_type=extension,
            document_type=document_type,
            content_hash=content_hash,
            text=full_text,
            pages=pages,
            metadata=document_metadata,
            entities=entities,
            processed_at=datetime.now(
                timezone.utc
            ),
        )

    # =====================================================
    # PROCESS FILE
    # =====================================================

    def process_file(
        self,
        path: str | Path,
        *,
        company_id: str,
        document_id: str | None = None,
        metadata: Mapping[str, Any] | None = None,
    ) -> ProcessedDocument:
        path = Path(path)

        if not path.is_file():
            raise DocumentValidationError(
                "Document file does not exist"
            )

        max_bytes = (
            self.config.max_file_size_mb
            * 1024
            * 1024
        )

        with path.open("rb") as stream:
            data = stream.read(max_bytes + 1)

        return self.process_bytes(
            data=data,
            filename=path.name,
            company_id=company_id,
            document_id=document_id,
            metadata=metadata,
        )

    # =====================================================
    # ASYNC PROCESSING
    # =====================================================

    async def process_file_async(
        self,
        path: str | Path,
        *,
        company_id: str,
        document_id: str | None = None,
        metadata: Mapping[str, Any] | None = None,
    ) -> ProcessedDocument:
        return await asyncio.to_thread(
            self.process_file,
            path,
            company_id=company_id,
            document_id=document_id,
            metadata=metadata,
        )

    # =====================================================
    # BATCH PROCESSING
    # =====================================================

    async def process_batch(
        self,
        paths: Sequence[str | Path],
        *,
        company_id: str,
    ) -> list[ProcessedDocument]:
        results = []

        for path in paths:
            document = await self.process_file_async(
                path,
                company_id=company_id,
            )

            results.append(document)

        return results


# =========================================================
# EXAMPLE
# =========================================================

if __name__ == "__main__":
    processor = DocumentProcessor(
        DocumentProcessingConfig(
            enable_ocr=False,
            max_file_size_mb=30,
            max_pages=300,
        )
    )

    sample = (
        "Motor M-101 Maintenance Report\n\n"
        "The motor temperature reached 95°C. "
        "Vibration increased to 12.5 mm/s. "
        "Possible causes include bearing wear "
        "and insufficient lubrication.\n\n"
        "Follow the approved safety procedure "
        "before maintenance."
    ).encode("utf-8")

    result = processor.process_bytes(
        data=sample,
        filename="motor_report.txt",
        company_id="company-001",
        metadata={
            "asset_id": "M-101",
            "site_id": "plant-01",
        },
    )

    print(
        json.dumps(
            result.to_dict(),
            indent=2,
            default=str,
        )
    )
